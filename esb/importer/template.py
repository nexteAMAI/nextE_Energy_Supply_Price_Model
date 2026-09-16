"""Writer for the standard upload template (D-TPL) and for filled fixture workbooks.

`build_template(...)` writes a blank template for one input class; `write_delivery(...)` writes
a filled workbook (used for the Reference Case fixtures and for tests). Both produce exactly
the layout that esb.importer.reader expects.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from esb import grid
from esb.importer.contract import (
    CONTROL_KEYS,
    RAW_FIXED_COLUMNS,
    TEMPLATE_VERSION,
    Control,
    Registry,
    preset_registry,
)

# Styling as measured on the CEO's canonical input templates (00_Spec/01_current/input_data_template, 16.09.2026, D113):
# Montserrat 11 bold white on navy for the header row, Montserrat 10 body, input cells on #FFFCEB (brand xlsx
# input fill), derived / locked cells on #EEF3FA, notes in muted grey, hairline borders #C8C8C6, gridlines off,
# freeze at A2, unit-bearing number formats with the unit outside the quotes.
NAVY = "1F3E66"  # nexte-brand navy (app.brand.NAVY)
INK = "0E1C2E"
MUTED = "6A6A6A"
LINE = "C8C8C6"
INPUT_YELLOW = "FFFCEB"
DERIVED_BLUE = "EEF3FA"
GREY = INPUT_YELLOW  # kept for callers
HEAD_FONT = Font(name="Montserrat", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Montserrat", size=10, color="000000")
BODY_BOLD = Font(name="Montserrat", size=10, bold=True, color="000000")
NOTE_FONT = Font(name="Montserrat", size=10, color=MUTED)
TITLE_FONT = Font(name="Montserrat", size=16, bold=True, color=NAVY)
EYEBROW_FONT = Font(name="Montserrat", size=11, bold=True, color=NAVY)
HEADING_FONT = Font(name="Montserrat", size=10, bold=True, color=NAVY)
DERIVED_FILL = PatternFill("solid", fgColor=DERIVED_BLUE)
HAIR = Side(style="thin", color=LINE)
BORDER = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
NUM_FORMAT = r'#,##0.0000\ "MWh";\(#,##0.0000\ "MWh"\)'  # quarter-hour energy, unit outside the quotes
UNIT_FORMATS = {
    "MWh": r'#,##0.0000\ "MWh";\(#,##0.0000\ "MWh"\)',
    "EUR/MWh": r'#,##0.00\ "€/MWh";\(#,##0.00\ "€/MWh"\);"-"',
    "ratio": "0.0000",
    "flag": "@",
    "MW": r'#,##0.000\ "MW";\(#,##0.000\ "MW"\)',
}
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor=GREY)

INSTRUCTIONS = [
    "nextE Energy Supply Bid Management Tool - standard quarter-hourly upload template",
    f"Template version {TEMPLATE_VERSION}. Sheets: Std_Control, Series_Registry, RAW_EET_QH, Recon_Check.",
    "",
    "01 What to deliver",
    "One workbook per input class and spine year. Paste values only into RAW_EET_QH (no formulas, no links).",
    "Columns A:C carry the local date (EET, Romanian clock), the interval start and the interval end.",
    "Columns D onward are the series slots E01.. (numbers) and C01.. (categories) exactly as declared in Series_Registry.",
    "",
    "02 Time basis (Std_Control!time_basis)",
    "local_clock: rows follow the real local clock. The last Sunday of March has 92 rows (03:00-03:45 absent);",
    "  the last Sunday of October has 100 rows (03:00-03:45 twice). The importer injects / merges by the declared rules.",
    "fixed_96: rows are already 96 per day for every day (positional grid). No daylight-saving rule is applied.",
    "",
    "03 Series registry",
    "Every slot with data must be declared: name, unit, class (Extensive = energy per interval; Intensive = price, ratio, power;",
    "Categorical = flag), the spring and autumn rules, and for volume series the sign convention k (-1 DSO metering,",
    "+1 BRP imbalance). Declared once per series, never inferred from the values.",
    "",
    "04 Rules applied at import (local_clock only)",
    "Extensive: injected interval = 0; merged interval = sum. Intensive: injected = linear interpolation or last value",
    "(spring_rule); merged = mean or volume-weighted mean over paired_volume_slot (autumn_rule). Categorical: carry forward / first.",
    "",
    "05 What the importer refuses",
    "Formulas in RAW_EET_QH; undeclared columns with data; blanks inside a required series (gaps are never filled);",
    "missing or duplicated intervals other than the daylight-saving pattern; dates outside the spine year;",
    "a volume series without k; a registry that contradicts itself.",
    "",
    "06 Formats",
    "Dates dd.mm.yyyy, times hh:mm. Decimal separator follows your Excel locale; the stored values are what count.",
    "Off-takers are referenced by code (OT1, OT2, ...); the name-to-code mapping is kept in the application, not in this file.",
]


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = CENTER if c > 1 else LEFT
        cell.border = BORDER
    ws.row_dimensions[1].height = 25
    ws.sheet_format.defaultRowHeight = 25
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    _style_body(ws)


def _style_body(ws) -> None:
    """Montserrat on every written cell below the header (brand: one typeface on every surface); hairline borders."""
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.value is not None or cell.fill.fgColor.rgb not in (None, "00000000"):
                if cell.font.name != "Montserrat":
                    cell.font = BODY_FONT
                cell.border = BORDER


def _write_instructions(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instructions"
    ws["A1"] = "CONFIDENTIAL - nextE"
    ws["A1"].font = EYEBROW_FONT
    for i, line in enumerate(INSTRUCTIONS, start=2):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = TITLE_FONT if i == 2 else (HEADING_FONT if line[:2].isdigit() else NOTE_FONT)
    for r in range(1, 5):
        ws.row_dimensions[r].height = 25
    ws.column_dimensions["A"].width = 120
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def _write_control(wb: Workbook, control: Control) -> None:
    ws = wb.create_sheet("Std_Control")
    ws.append(["key", "value", "note"])
    notes = {
        "template_version": "do not change",
        "input_class": "offtaker_load | pv_generation | baseload_nomination | wholesale_prices",
        "spine_year": "calendar year of every row in RAW_EET_QH",
        "intervals_per_day": "96 (quarter-hourly)",
        "time_basis": "local_clock | fixed_96 (see Instructions 02)",
        "offset_eet_cet_h": "1 (fixed)",
        "scenario": "free text, e.g. Aurora Central Q2/2026 (wholesale_prices only)",
        "provider": "who produced the data",
        "delivery_date": "dd.mm.yyyy",
        "notes": "free text",
    }
    values = {
        "template_version": control.template_version,
        "input_class": control.input_class,
        "spine_year": control.spine_year,
        "intervals_per_day": control.intervals_per_day,
        "time_basis": control.time_basis,
        "offset_eet_cet_h": control.offset_eet_cet_h,
        "scenario": control.scenario,
        "provider": control.provider,
        "delivery_date": control.delivery_date,
        "notes": control.notes,
    }
    for k in CONTROL_KEYS:
        ws.append([k, values[k], notes[k]])
        r = ws.max_row
        v = ws.cell(row=r, column=2)
        v.fill = INPUT_FILL
        v.alignment = CENTER
        v.number_format = "0" if k in ("spine_year", "intervals_per_day", "offset_eet_cet_h") else ('dd"."mm"."yyyy' if k == "delivery_date" else "@")
        ws.cell(row=r, column=1).font = BODY_FONT
        ws.cell(row=r, column=3).font = NOTE_FONT
    _style_header(ws, 3)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 70


def registry_columns(input_class: str) -> list[str]:
    """Column order of Series_Registry per input class (1.1, as on the CEO's canonical templates): the wholesale
    file carries scenario_name after slot; the entity files carry entity_code and entity_name after slot."""
    if input_class == "wholesale_prices":
        return ["slot", "scenario_name", "name", "unit", "class", "spring_rule", "autumn_rule", "paired_volume_slot", "k", "basis",
                "entity_code", "required", "notes"]
    return ["slot", "entity_code", "entity_name", "name", "unit", "class", "spring_rule", "autumn_rule", "paired_volume_slot", "k", "basis",
            "required", "notes"]


def _write_registry(wb: Workbook, registry: Registry, input_class: str = "offtaker_load") -> None:
    ws = wb.create_sheet("Series_Registry")
    cols = registry_columns(input_class)
    ws.append(cols)
    editable = {"scenario_name", "entity_name", "entity_code", "name"}  # the deliverer names the block / entity; the rest is locked
    for s in registry.slots:
        vals = {
            "slot": s.slot, "scenario_name": s.scenario_name or "", "entity_code": s.entity_code or "",
            "entity_name": s.entity_name or (f"entity_name_{s.entity_code}" if s.entity_code else ""), "name": s.name, "unit": s.unit,
            "class": s.cls, "spring_rule": s.spring_rule, "autumn_rule": s.autumn_rule, "paired_volume_slot": s.paired_volume_slot or "",
            "k": s.k if s.k is not None else "", "basis": s.basis or "", "required": "Y" if s.required else "N", "notes": s.notes,
        }
        ws.append([vals[c] for c in cols])
        r = ws.max_row
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=j)
            cell.fill = INPUT_FILL if c in editable else DERIVED_FILL
            cell.font = BODY_BOLD if c == "slot" else (NOTE_FONT if c == "notes" else BODY_FONT)
            cell.alignment = LEFT if c in ("name", "notes", "scenario_name") else CENTER
            cell.number_format = "0" if c == "k" else "@"
    _style_header(ws, len(cols))
    widths = {"slot": 9, "scenario_name": 30, "entity_code": 12, "entity_name": 30, "name": 44, "unit": 10, "class": 12, "spring_rule": 14,
              "autumn_rule": 16, "paired_volume_slot": 18, "k": 6, "basis": 12, "required": 9, "notes": 80}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[c]


def _write_raw(wb: Workbook, registry: Registry, data: pd.DataFrame | None, year: int | None = None) -> int:
    ws = wb.create_sheet("RAW_EET_QH")
    slots = [s.slot for s in registry.slots]
    ws.append(list(RAW_FIXED_COLUMNS) + slots)
    n = 0
    if data is not None:
        cols = ["Date_EET", "Start_EET", "End_EET"] + [
            s.frame_name if s.frame_name in data.columns else s.name for s in registry.slots  # 1.1: scenario blocks keyed by frame name
        ]
        missing = [c for c in cols if c not in data.columns]
        if missing:
            raise ValueError(f"data lacks columns {missing}")
        for row in data[cols].itertuples(index=False):
            vals = list(row)
            out = [vals[0], vals[1], vals[2]] + [
                (None if (isinstance(v, float) and v != v) else v) for v in vals[3:]
            ]
            ws.append(out)
            n += 1
    _style_header(ws, 3 + len(slots))
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    fmt_of = {s.slot: ("@" if s.cls == "Categorical" else UNIT_FORMATS.get(s.unit, NUM_FORMAT)) for s in registry.slots}
    for j, s in enumerate(registry.slots, start=4):  # column-level formats so typed values render 1.234,5000 MWh in a Romanian Excel
        col = ws.column_dimensions[get_column_letter(j)]
        col.width = 25.6
        col.number_format = fmt_of[s.slot]
        col.font = BODY_FONT
        col.fill = INPUT_FILL
    for letter, fmt in (("A", 'dd"."mm"."yyyy'), ("B", "hh:mm"), ("C", "hh:mm")):
        ws.column_dimensions[letter].number_format = fmt
        ws.column_dimensions[letter].font = BODY_FONT
        ws.column_dimensions[letter].fill = INPUT_FILL
    # the frame of the spine year is pre-styled row by row (the canonical template's paste surface: yellow input cells,
    # hairline borders, unit-bearing formats), so a delivery pasted over it keeps the look
    last = max(ws.max_row, 1 + grid.grid_info(int(year)).rows) if year else ws.max_row
    for r in range(2, last + 1):
        for j in range(1, 4 + len(slots)):
            cell = ws.cell(row=r, column=j)
            cell.fill = INPUT_FILL
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER
            cell.number_format = 'dd"."mm"."yyyy' if j == 1 else ("hh:mm" if j in (2, 3) else fmt_of[slots[j - 4]])
    ws.sheet_view.topLeftCell = "A2"
    return n


def _write_recon(wb: Workbook, control: Control) -> None:
    ws = wb.create_sheet("Recon_Check")
    g = grid.grid_info(int(control.spine_year))
    required = g.rows  # local_clock delivers -4 +4 rows, the same count
    ws.append(["check", "value", "required", "status"])
    ws.append(["rows delivered (dates in RAW_EET_QH)", "=COUNT(RAW_EET_QH!A:A)", required, '=IF(B2=C2,"PASS","FAIL")'])
    ws.append(["first date", "=MIN(RAW_EET_QH!A:A)", date(g.year, 1, 1), '=IF(B3=C3,"PASS","FAIL")'])
    ws.append(["last date", "=MAX(RAW_EET_QH!A:A)", date(g.year, 12, 31), '=IF(B4=C4,"PASS","FAIL")'])
    ws.append(["note", "The binding checks (six) run in the application at upload; this sheet is a convenience.", "", ""])
    for r in range(2, 5):
        ws.cell(row=r, column=1).font = BODY_FONT
        ws.cell(row=r, column=2).fill = DERIVED_FILL
        ws.cell(row=r, column=3).fill = INPUT_FILL
        ws.cell(row=r, column=4).fill = DERIVED_FILL
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).alignment = CENTER
        ws.cell(row=r, column=2).number_format = "0" if r == 2 else 'dd"."mm"."yyyy;;"-"'
        ws.cell(row=r, column=3).number_format = "0" if r == 2 else 'dd"."mm"."yyyy'
        ws.cell(row=r, column=4).number_format = "@"
    ws.cell(row=5, column=2).font = NOTE_FONT
    _style_header(ws, 4)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 14


def write_delivery(
    path: str | Path,
    control: Control,
    registry: Registry,
    data: pd.DataFrame | None = None,
) -> Path:
    """Write a template (data=None) or a filled delivery.

    `data` columns: Date_EET (date/datetime), Start_EET (time), End_EET (time) and one column per
    registry slot *name*. NaN is written as an empty cell.
    """
    errs = control.validate() + registry.validate()
    if errs:
        raise ValueError("; ".join(errs))
    wb = Workbook()
    _write_instructions(wb)
    _write_control(wb, control)
    _write_registry(wb, registry, control.input_class)
    _write_raw(wb, registry, data, year=int(control.spine_year) if data is None else None)
    _write_recon(wb, control)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build_template(
    input_class: str,
    spine_year: int,
    path: str | Path,
    entity_codes: list[str] | None = None,
    time_basis: str = "local_clock",
    scenarios: list[str] | None = None,
) -> Path:
    """A blank canonical template (contract 1.1) for the entities of the register: off-taker codes, the PV plant(s),
    or the scenario blocks for a wholesale file; the styling of the CEO's canonical templates (D113)."""
    control = Control(input_class=input_class, spine_year=spine_year, time_basis=time_basis)
    return write_delivery(path, control, preset_registry(input_class, entity_codes, scenarios))


def raw_frame_fixed_96(year: int, series: dict[str, pd.Series | list]) -> pd.DataFrame:
    """Helper: build a fixed_96 RAW frame (days x 96 rows) from per-row series keyed by slot name."""
    g = grid.make_grid(year)
    out = pd.DataFrame(
        {
            "Date_EET": [d.date() for d in g["date"]],
            "Start_EET": [t.time() for t in g["start_eet"]],
            "End_EET": [t.time() for t in g["end_eet"]],
        }
    )
    for name, s in series.items():
        out[name] = list(s)
    return out
