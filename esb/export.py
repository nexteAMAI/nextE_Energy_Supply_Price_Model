"""esb.export - Excel and CSV exports of a run (execution prompt section 10, page 11).

The workbook mirrors the five result sheets of the Reference Case, one sheet per off-taker for
the price build-up, the parameter register, the QH result aggregated by day (the full
quarter-hour frame on request) and the provenance of the run. Display names are allowed in
exports (F-035); every sheet carries the confidentiality marking. Number formats use Excel's
locale-aware tokens so a Romanian Excel renders 84,50 and 1.490.
"""

from __future__ import annotations

import io
import json
from datetime import UTC, date, datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from esb import __version__
from esb.engine import RunResult
from esb.labels import RESELL, RETAIL, TOTAL, label

NAVY = "1F3E66"
INK = "0E1C2E"
PAPER = "F7F6F3"
LINE = "C8C8C6"
FONT = "Montserrat"
CONFIDENTIAL = "CONFIDENTIAL - nextE"
MONTH_COLUMNS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Year"]
NUM = "#,##0.00;(#,##0.00);-"
NUM0 = "#,##0;(#,##0);-"
PCT = "0.0%;(0.0%);-"

_hdr_font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
_hdr_fill = PatternFill("solid", fgColor=NAVY)
_body_font = Font(name=FONT, size=9, color=INK)
_title_font = Font(name=FONT, size=12, bold=True, color=NAVY)
_eyebrow_font = Font(name=FONT, size=8, bold=True, color="6A6A6A")
_alt_fill = PatternFill("solid", fgColor=PAPER)
_hair = Side(style="hair", color=LINE)


def _sheet_header(ws, title: str, caption: str, run: RunResult) -> int:
    ws["A1"] = CONFIDENTIAL
    ws["A1"].font = _eyebrow_font
    ws["A2"] = title
    ws["A2"].font = _title_font
    ws["A3"] = caption
    ws["A3"].font = _body_font
    ws["A4"] = (f"Engine v{__version__} · scenario {run.params.scenario_active} · year {run.params.spine_year} · "
                f"exported {datetime.now(UTC).strftime('%d.%m.%Y %H:%M')} UTC · EUR unless stated")
    ws["A4"].font = _eyebrow_font
    return 6


def _write_table(ws, row0: int, df: pd.DataFrame, index_label: str = "", pct_rows: set[str] | None = None, int_cols: set[str] | None = None) -> int:
    pct_rows = pct_rows or set()
    int_cols = int_cols or set()
    cols = [index_label, *[str(c) for c in df.columns]]
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row0, column=j, value=c)
        cell.font, cell.fill = _hdr_font, _hdr_fill
        cell.alignment = Alignment(horizontal="left" if j == 1 else "right", vertical="center")
        cell.border = Border(bottom=_hair)
    r = row0 + 1
    for i, (idx, rec) in enumerate(df.iterrows()):
        if isinstance(idx, (pd.Timestamp, datetime, date)):  # dates as dates (dd.mm.yyyy), not as text
            c0 = ws.cell(row=r, column=1, value=idx.to_pydatetime().date() if isinstance(idx, pd.Timestamp) else idx)
            c0.number_format = "dd.mm.yyyy"
        else:
            c0 = ws.cell(row=r, column=1, value=str(idx))
        c0.font = _body_font
        if i % 2 == 1:
            c0.fill = _alt_fill
        for j, (name, v) in enumerate(rec.items(), start=2):
            if isinstance(v, (float, np.floating)):
                v = None if np.isnan(v) else float(v)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (pd.Timestamp, datetime)):
                v = v.to_pydatetime() if isinstance(v, pd.Timestamp) else v
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = _body_font
            cell.alignment = Alignment(horizontal="right")
            if isinstance(v, (int, float)):
                if str(idx) in pct_rows or str(idx).endswith("_pct"):
                    cell.number_format = PCT
                elif str(name) in int_cols:
                    cell.number_format = NUM0
                else:
                    cell.number_format = NUM
            if i % 2 == 1:
                cell.fill = _alt_fill
            cell.border = Border(bottom=_hair)
        r += 1
    ws.column_dimensions["A"].width = 42
    for j in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = ws.cell(row=row0 + 1, column=2)
    return r + 1


def _labelled(df: pd.DataFrame, **kw) -> pd.DataFrame:
    out = df.copy()
    out.index = [label(k, **kw) for k in out.index]
    return out


def _sections_frame(run: RunResult) -> dict[str, pd.DataFrame]:
    return {run.params.offtaker(code).label: T.frame() for code, T in run.pnl.sections.items()}


def build_workbook(run: RunResult, include_qh: bool = False, scenario_name: str = "", provenance: list[dict] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    P = run.params

    # ---- Portf Overview ---------------------------------------------------------------------
    ws = wb.create_sheet("Portf Overview")
    r = _sheet_header(ws, "Portfolio overview", "Year values by leg; budget vs forecast; sections 5-7 below", run)
    ov = run.overview.table.copy()
    ren = {}
    for o in P.offtakers:
        ren[f"{o.code}_budget"] = f"{o.label} budget"
        ren[f"{o.code}_forecast"] = f"{o.label} forecast"
    ov = ov.rename(columns={"portfolio_budget": "Portfolio budget", "portfolio_forecast": "Portfolio forecast", "delta": "Delta", "resell": RESELL, **ren})
    r = _write_table(ws, r, _labelled(ov), "Line")
    for title, block in (("Cash flow", run.overview.cashflow), ("Pricing (selected off-taker)", run.overview.pricing), ("Checks and tripwires", run.overview.checks)):
        ws.cell(row=r, column=1, value=title).font = _title_font
        r += 1
        df = pd.DataFrame({"Value": list(block.values())}, index=[label(k, leg=(TOTAL if title == "Cash flow" else RETAIL if title.startswith("Pricing") else "")) for k in block])
        r = _write_table(ws, r, df, "Item")

    # ---- Cons_P&L ---------------------------------------------------------------------------
    ws = wb.create_sheet("Cons_P&L")
    r = _sheet_header(ws, "Consolidated P&L", "Monthly and year; portfolio block then one block per off-taker (EUR, MWh, EUR/MWh)", run)
    pf = run.pnl.portfolio.frame()
    pf.columns = MONTH_COLUMNS
    r = _write_table(ws, r, _labelled(pf), "Portfolio")
    for name, T in _sections_frame(run).items():
        ws.cell(row=r, column=1, value=name).font = _title_font
        r += 1
        T.columns = MONTH_COLUMNS
        r = _write_table(ws, r, _labelled(T, leg=RETAIL), name)

    # ---- CF_Mth -----------------------------------------------------------------------------
    ws = wb.create_sheet("CF_Mth")
    r = _sheet_header(ws, "Monthly cash flow", "Accruals, settlement keys, receipts, payments, VAT, financing (EUR)", run)
    cf = run.cashflow.frame()
    cf.columns = [*MONTH_COLUMNS[:12], "Beyond Dec", "Year"]
    r = _write_table(ws, r, _labelled(cf, fallback=TOTAL), "Line")

    # ---- CF_Daily_Ledger --------------------------------------------------------------------
    ws = wb.create_sheet("CF_Daily_Ledger")
    r = _sheet_header(ws, "Daily cash ledger", "365 settlement days; floor, injection, loan, interest (EUR)", run)
    if run.cashflow.daily is not None:
        d = run.cashflow.daily.copy()
        if "date" in d.columns:
            d = d.set_index("date")
        d.columns = [label(c, fallback=TOTAL) for c in d.columns]
        r = _write_table(ws, r, d, "Date")
    ws.cell(row=r, column=1, value="Summary").font = _title_font
    r += 1
    r = _write_table(ws, r, pd.DataFrame({"Value": list(run.cashflow.daily_summary.values())}, index=[label(k, fallback=TOTAL) for k in run.cashflow.daily_summary]), "Item")

    # ---- Pricing_Calc per off-taker -------------------------------------------------------------
    for code, pr in run.pricing.items():
        o = P.offtaker(code)
        ws = wb.create_sheet(f"Pricing_{code}"[:31])
        r = _sheet_header(ws, f"Price build-up - {o.label}", "Year column and months (EUR/MWh unless stated); corrected logic of D87/D88" if not pr.as_cached else "Workbook-cached logic", run)
        rows = {}
        for k, v in pr.year.items():
            months = pr.months.get(k)
            rows[k] = [v, *(list(months) if months is not None else [np.nan] * 12)]
        df = pd.DataFrame(rows, index=["Year", *MONTH_COLUMNS[:12]]).T
        r = _write_table(ws, r, _labelled(df, leg=RETAIL), "Line")
        if pr.manual:
            ws.cell(row=r, column=1, value="Manual case").font = _title_font
            r += 1
            r = _write_table(ws, r, pd.DataFrame({"Value": list(pr.manual.values())}, index=[label(k, leg=RETAIL) for k in pr.manual]), "Line")

    # ---- Guarantees ----------------------------------------------------------------------------
    ws = wb.create_sheet("Guarantees")
    r = _sheet_header(ws, "Guarantees", "Outstanding per counterparty and off-taker by month (EUR) and BGL fees", run)
    keys = [k for k in run.pnl.portfolio.order if k.startswith("g_out_") or k.startswith("g_fee_") or k in ("guarantees_outstanding", "market_bgl_fees")]
    gf = pd.DataFrame({k: run.pnl.portfolio[k] for k in keys}, index=MONTH_COLUMNS).T
    r = _write_table(ws, r, _labelled(gf), "Line")
    own = pd.DataFrame({f"{P.offtaker(c).label}": T["own_guarantee"] for c, T in run.pnl.sections.items() if "own_guarantee" in T}, index=MONTH_COLUMNS).T
    if len(own):
        ws.cell(row=r, column=1, value="Own guarantees (issued to off-takers)").font = _title_font
        r += 1
        r = _write_table(ws, r, own, "Off-taker")

    # ---- QH ----------------------------------------------------------------------------------
    ws = wb.create_sheet("QH_daily")
    r = _sheet_header(ws, "Quarter-hour engine - daily aggregates", "Sums per day of the main QH columns (MWh, EUR); prices as daily means", run)
    q = run.qh.qh
    g = run.grid
    daily = q.groupby(g["date"].dt.date.values).agg({c: ("mean" if c in ("dam", "idct", "surplus_price", "deficit_price") else "sum")
                                                    for c in q.columns if q[c].dtype.kind == "f"})
    daily.columns = [label(c) for c in daily.columns]
    r = _write_table(ws, r, daily, "Date")
    if include_qh:
        ws = wb.create_sheet("QH_full")
        r = _sheet_header(ws, "Quarter-hour engine - full frame", "35.040 rows; every named column of the engine", run)
        full = q.drop(columns=[c for c in ("date", "interval") if c in q.columns]).copy()
        full.insert(0, "interval", g["interval"].values)
        full.insert(0, "date", g["date"].dt.date.values)
        full = full.set_index("date")
        r = _write_table(ws, r, full, "Date", int_cols={"interval"})

    # ---- Parameters ----------------------------------------------------------------------------
    ws = wb.create_sheet("Parameters")
    r = _sheet_header(ws, "Parameter register", f"Scenario file: {scenario_name or '(unsaved)'}", run)
    flat = _flatten(P.to_dict())
    df = pd.DataFrame({"Value": [str(v) for v in flat.values()]}, index=list(flat.keys()))
    r = _write_table(ws, r, df, "Parameter")

    # ---- Provenance --------------------------------------------------------------------------
    ws = wb.create_sheet("Provenance")
    r = _sheet_header(ws, "Provenance", "Inputs of this run and the calculation-order trace", run)
    prov = provenance or []
    if prov:
        df = pd.DataFrame(prov).astype(str)
        df.index = [str(i + 1) for i in range(len(df))]
        r = _write_table(ws, r, df, "#")
    tr = pd.DataFrame({"Cumulative s": [t for _, t in run.trace]}, index=[n for n, _ in run.trace])
    ws.cell(row=r, column=1, value="Calculation-order trace (G0-D6)").font = _title_font
    r += 1
    r = _write_table(ws, r, tr, "Stage")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _flatten(d: dict, prefix: str = "") -> dict[str, object]:
    out: dict[str, object] = {}
    for k, v in d.items():
        key = f"{prefix}.{k}" if prefix else str(k)
        if isinstance(v, dict):
            out.update(_flatten(v, key))
        elif isinstance(v, list) and v and isinstance(v[0], dict):
            for i, item in enumerate(v):
                tag = item.get("code", str(i)) if isinstance(item, dict) else str(i)
                out.update(_flatten(item, f"{key}[{tag}]"))
        else:
            out[key] = json.dumps(v) if isinstance(v, list) else v
    return out


def csv_bytes(df: pd.DataFrame, index: bool = True) -> bytes:
    """CSV with Romanian conventions: semicolon separator, decimal comma."""
    return df.to_csv(sep=";", decimal=",", index=index, float_format="%.6f").encode("utf-8-sig")
