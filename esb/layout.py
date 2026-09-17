"""esb.layout - the canonical output workbook (D113, TPL-CANON).

The workbook is rendered from data/layout/run_export.json, the layout specification extracted from the CEO's
formatted template (tools/extract_layout.py), with the values of a run. The specification carries, per sheet,
the metadata band, the header row, the column widths, every row (engine key, role, unit) and the block rhythm
(blank separators, section headers); the standard (nexte-output-workbook-formatting.md) supplies the palette,
the type scale and the number-format catalogue, which live here as code. Nothing of the CEO's template files
enters the repository: the specification is a few hundred KB of structure, the values come from the engine.

Rules applied (audit rulings D-C .. D-H): one off-taker block template replicated per off-taker with the display
name in the section header; units from the engine (esb.labels.unit_of), never from the template; decimals by
role (monthly 0 dp volumes, daily 3 dp, quarter-hour 4 dp; prices 2 dp, KPI price rows 3 dp; checks 6 dp);
no merged cells (Center Across Selection); no formulas; gridlines off; hairline borders only on cells with
content or fill; A4 of every sheet is the application's stamp (engine version, scenario, year, export time).
"""

from __future__ import annotations

import io
import json
from datetime import UTC, date, datetime
from functools import cache, lru_cache
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.indexed_list import IndexedList

from esb import __version__
from esb.calendar_ro import CALENDAR_COLUMNS, calendar_block, holiday_source
from esb.engine import RunResult
from esb.export import _flatten
from esb.labels import RESELL, RETAIL, TOTAL, label, unit_of

SPEC_PATH = Path(__file__).resolve().parents[1] / "data" / "layout" / "run_export.json"
FONT = "Montserrat"
NAVY, INK, MUTED, LINE, WHITE, BLACK = "1F3E66", "0E1C2E", "6A6A6A", "C8C8C6", "FFFFFF", "000000"
DATA_FILL, GROUP_FILL, INDEX_FILL = "EEF3FA", "F4F4F4", "F7F6F3"
CONFIDENTIAL = "CONFIDENTIAL - nextE"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ROW_H, ROW_H_WRAP, ROW_H_KEY = 25.0, 100.0, 20.25

# ---- style catalogue (the standard, as code) --------------------------------------------------------------
_hair = Side(style="hair", color=LINE)
BORDER = Border(left=_hair, right=_hair, top=_hair, bottom=_hair)
NO_BORDER = Border()
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
CAS = Alignment(horizontal="centerContinuous", vertical="center")  # Center Across Selection - no merges


@cache
def font(size: float = 12, bold: bool = False, italic: bool = False, color: str = BLACK) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


@cache
def fill(rgb: str | None) -> PatternFill | None:
    return PatternFill("solid", fgColor=rgb) if rgb else None


ROLE_STYLE = {  # role -> (fill of value cells, font)
    "data": (DATA_FILL, font()),
    "subtotal": (DATA_FILL, font(bold=True)),
    "total": (INK, font(bold=True, color=WHITE)),
    "check": (DATA_FILL, font()),
    "memo": (DATA_FILL, font(italic=True, color=MUTED)),
    "group": (GROUP_FILL, font(bold=True, color=MUTED)),
    "section": (NAVY, font(bold=True, color=WHITE)),
    "header": (INK, font(bold=True, color=WHITE)),
}

# number formats: unit -> {granularity: format}; the space before the unit is escaped as Excel stores it
def _nf(dp: int, unit: str, dash: bool = True) -> str:
    n = "#,##0" + ("." + "0" * dp if dp else "")
    core = f'{n}\\ "{unit}";\\({n}\\ "{unit}"\\)'
    return core + (';"-"' if dash else "")


NUMBER_FORMATS: dict[str, dict[str, str]] = {
    "EUR": {"monthly": _nf(0, "€"), "daily": _nf(0, "€"), "qh": _nf(2, "€"), "kpi": _nf(0, "€")},
    "EUR/MWh": {"monthly": _nf(2, "€/MWh"), "daily": _nf(2, "€/MWh"), "qh": _nf(2, "€/MWh"), "kpi": _nf(3, "€/MWh")},
    "MWh": {"monthly": _nf(0, "MWh", False), "daily": _nf(3, "MWh", False), "qh": _nf(4, "MWh", False), "kpi": _nf(0, "MWh", False)},
    "MW": {"monthly": _nf(1, "MW", False), "daily": _nf(1, "MW", False), "qh": _nf(3, "MW", False), "kpi": _nf(1, "MW", False)},
    "%": {"*": '0.0%;\\(0.0%\\);"-"'},
    "ratio": {"*": "0.0000"},
    "check": {"*": "0.000000"},
    "GC": {"*": _nf(0, "GC", False)},
    "GC/MWh": {"*": '#,##0.000\\ "GC/MWh"'},
    "RON/GC": {"*": '#,##0.0000\\ "RON/GC"'},
    "RON/EUR": {"*": '#,##0.0000\\ "RON/€"'},
    "EUR/GC": {"*": '#,##0.00\\ "€/GC"'},
    "#": {"*": "0"}, "days": {"*": "0"}, "sign": {"*": "0"}, "flag": {"*": "@"}, "s": {"*": "0.000"}, "": {"*": "@"},
    "date": {"*": 'dd"."mm"."yyyy'}, "month": {"*": "mmm\\ yyyy"}, "time": {"*": "hh:mm"}, "int": {"*": "0"}, "text": {"*": "@"},
}


def number_format(unit: str | None, granularity: str = "monthly") -> str:
    table = NUMBER_FORMATS.get(unit or "", NUMBER_FORMATS["EUR"])
    return table.get(granularity) or table.get("*") or table["monthly"]


def unit_text(unit: str | None) -> str:
    """Unit column text: the engine unit in the house spelling (EUR, EUR/MWh, MWh, MW, %, check ...)."""
    return {"": "", None: "", "ratio": "ratio", "#": "#", "sign": "sign", "flag": "flag"}.get(unit, unit or "")


@lru_cache(maxsize=1)
def spec() -> dict:
    with open(SPEC_PATH, encoding="utf-8") as f:
        return json.load(f)


# ---- sheet buffer -------------------------------------------------------------------------------------
class Sheet:
    """A sheet under construction: sparse cells keyed by (row, col), row heights, widths and the freeze pane.
    flush() writes it as a write-only worksheet (inline strings, streamed) so the wide grids and the small
    sheets share one code path and one workbook."""

    def __init__(self, name: str):
        self.name = name
        self.cells: dict[tuple[int, int], tuple] = {}
        self.heights: dict[int, float] = {}
        self.widths: dict[str, float] = {}
        self.freeze: str | None = None
        self.max_row = 0

    def put(self, r: int, c: int, value=None, *, f: Font | None = None, bg: str | None = None, nf: str | None = None,
            al: Alignment = CENTER, border: bool | None = None) -> None:
        if isinstance(value, (np.floating, np.integer)):
            value = value.item()
        if isinstance(value, float) and not np.isfinite(value):
            value = None
        elif isinstance(value, (pd.Timestamp, np.datetime64)):
            value = pd.Timestamp(value).to_pydatetime()
        if border is None:
            border = value is not None or bool(bg)
        self.cells[(r, c)] = (value, f or font(), bg, nf, al, border)
        self.max_row = max(self.max_row, r)

    def height(self, r: int, h: float = ROW_H) -> None:
        self.heights[r] = h
        self.max_row = max(self.max_row, r)

    def flush(self, wb: Workbook) -> None:
        ws = wb.create_sheet(self.name)
        ws.sheet_view.showGridLines = False
        for col, w in self.widths.items():
            ws.column_dimensions[col].width = w
        if self.freeze:
            ws.freeze_panes = self.freeze
        for r, h in self.heights.items():
            ws.row_dimensions[r].height = h
        rows: dict[int, dict[int, tuple]] = {}
        for (r, c), spec_cell in self.cells.items():
            rows.setdefault(r, {})[c] = spec_cell
        for r in range(1, self.max_row + 1):
            row = rows.get(r)
            if not row:
                ws.append([])
                continue
            width = max(row)
            out = []
            for c in range(1, width + 1):
                sc = row.get(c)
                if sc is None:
                    out.append(None)
                    continue
                value, f, bg, nf, al, border = sc
                cell = WriteOnlyCell(ws)
                cell.font = f
                if bg:
                    cell.fill = fill(bg)
                if nf:
                    cell.number_format = nf
                elif isinstance(value, (datetime, date)):
                    cell.number_format = NUMBER_FORMATS["date"]["*"]
                cell.alignment = al
                cell.border = BORDER if border else NO_BORDER
                cell.value = value
                out.append(cell)
            ws.append(out)


def _put(ws: Sheet, r: int, c: int, value=None, **kw) -> None:
    ws.put(r, c, value, **kw)


def _band(ws, title: str, subtitle: str, stamp: str, ncols: int) -> None:
    """Rows 1-4: banner, sheet title, subtitle, application stamp (A4 belongs to the application)."""
    texts = [(CONFIDENTIAL, font(11, True, color=NAVY)), (title, font(16, True, color=NAVY)), (subtitle, font(10, color=MUTED)), (stamp, font(10, color=MUTED))]
    for i, (t, f) in enumerate(texts, start=1):
        _put(ws, i, 1, t, f=f, al=LEFT, border=False)
        for c in range(2, ncols + 1):
            _put(ws, i, c, None, f=font(10, color=MUTED), border=False)
        ws.height(i)
    ws.height(5)


def _finish(ws: Sheet, widths: dict[str, float], freeze: str | None) -> None:
    ws.widths = dict(widths)
    ws.freeze = freeze


def _stamp(run: RunResult, scenario_name: str) -> str:
    now = datetime.now(UTC).strftime("%d.%m.%Y %H:%M UTC")
    return f"Engine v{__version__} · scenario {scenario_name or '(unsaved)'} · year {run.params.spine_year} · exported {now}"


# ---- label-column family ------------------------------------------------------------------------------
class LabelSheet:
    """A label-column sheet: A label, B unit, C spacer, then value columns with spacers between column groups.
    `columns` is the ordered list of (letter, header text, value source key) for the value columns."""

    def __init__(self, ws, columns: list[tuple[str, str]], widths: dict[str, float]):
        self.ws = ws
        self.columns = columns  # [(letter, header)]
        self.value_cols = [column_index_from_string(c) for c, _ in columns]
        self.spacers = sorted({c for c in range(3, max(self.value_cols) + 1) if c not in self.value_cols})
        self.ncols = max(self.value_cols)
        self.widths = widths

    def header(self, r: int, label_text: str, unit_text_: str = "Unit") -> None:
        bg, f = ROLE_STYLE["header"]
        _put(self.ws, r, 1, label_text, f=f, bg=bg, al=LEFT)
        _put(self.ws, r, 2, unit_text_, f=f, bg=bg)
        for (col, text) in self.columns:
            _put(self.ws, r, column_index_from_string(col), text, f=f, bg=bg)
        self.ws.height(r)

    def section(self, r: int, text: str) -> None:
        bg, f = ROLE_STYLE["section"]
        _put(self.ws, r, 1, text, f=f, bg=bg, al=LEFT)
        _put(self.ws, r, 2, None, f=f, bg=bg)
        for c in self.value_cols:
            _put(self.ws, r, c, None, f=f, bg=bg)
        self.ws.height(r)

    def group(self, r: int, text: str) -> None:
        bg, f = ROLE_STYLE["group"]
        _put(self.ws, r, 1, text, f=f, bg=bg, al=LEFT)
        self.ws.height(r)

    def blank(self, r: int) -> None:
        self.ws.height(r)

    def line(self, r: int, text: str, unit: str | None, role: str, values: list, granularity: str = "monthly") -> None:
        bg, f = ROLE_STYLE.get(role, ROLE_STYLE["data"])
        label_font = font(bold=role in ("total", "subtotal"), italic=role == "memo", color=MUTED if role == "memo" else BLACK)
        _put(self.ws, r, 1, text, f=label_font, al=LEFT)
        _put(self.ws, r, 2, unit_text(unit), f=label_font)
        nf = number_format(unit, granularity)
        for c, v in zip(self.value_cols, values, strict=False):
            _put(self.ws, r, c, v, f=f, bg=bg, nf=nf)
        self.ws.height(r)


def _month_columns(has_year: bool = True, beyond: bool = False) -> list[tuple[str, str]]:
    cols = [("D", "Year")] if has_year else []
    cols += [(get_column_letter(6 + i), m) for i, m in enumerate(MONTHS)]
    if beyond:
        cols.append(("R", "Beyond Dec"))
    return cols


def _render_rows(sheet: LabelSheet, rows: list[dict], start: int, values_of, unit_ctx: str | None = None, granularity: str = "monthly",
                 section_text=None) -> int:
    """Write the spec rows of one block from row `start`; values_of(key) -> list aligned to the sheet's columns."""
    r = start
    for row in rows:
        kind = row["kind"]
        if kind == "blank":
            sheet.blank(r)
        elif kind == "section":
            sheet.section(r, section_text(row) if section_text else row["text"])
        elif kind == "header":
            sheet.header(r, section_text(row) if section_text else row["text"])
        elif kind == "line":
            key = row.get("key")
            if row.get("role") == "group" and key is None:
                sheet.group(r, row["label"])
            else:
                vals = values_of(key, row) if key else []
                unit = unit_of(key, unit_ctx) if key else row.get("unit")
                sheet.line(r, row["label"], unit, row.get("role", "data"), vals, granularity)
        r += 1
    return r


# ---- sheets -------------------------------------------------------------------------------------------
def _overview(wb: Workbook, run: RunResult, stamp: str) -> None:
    sp = spec()["sheets"]["Portf Overview"]
    P = run.params
    ws = Sheet("Portf Overview")
    table = run.overview.table
    cols: list[tuple[str, str]] = [("D", "Portfolio budget"), ("E", "Portfolio forecast"), ("F", "Delta")]
    keys = ["portfolio_budget", "portfolio_forecast", "delta"]
    c = 8  # H
    for o in P.offtakers:
        cols += [(get_column_letter(c), f"{o.label} budget"), (get_column_letter(c + 1), f"{o.label} forecast")]
        keys += [f"{o.code}_budget", f"{o.code}_forecast"]
        c += 3
    cols.append((get_column_letter(c), RESELL))
    keys.append("resell")
    widths = {"A": 90.71, "B": 15.71}
    for col, _ in cols:
        widths[col] = 30.71
    for col in range(3, c + 1):
        if get_column_letter(col) not in widths:
            widths[get_column_letter(col)] = 2.71
    _band(ws, sp["band"]["A2"], sp["band"]["A3"], stamp, c)
    sheet = LabelSheet(ws, cols, widths)
    sheet.header(6, "Line")

    def values_of(key, row):
        block = row.get("block")
        if block == "overview":
            return [float(table.loc[key, k]) if k in table.columns and key in table.index else None for k in keys]
        src = run.overview.cashflow if block == "overview_cf" else run.overview.checks
        return [src.get(key)] + [None] * (len(keys) - 1)

    _render_rows(sheet, sp["rows"], 7, values_of, unit_ctx="overview")
    _finish(ws, widths, "C7")
    ws.flush(wb)


def _cons_pnl(wb: Workbook, run: RunResult, stamp: str) -> None:
    sp = spec()["sheets"]["Cons_P&L"]
    P = run.params
    ws = Sheet("Cons_P&L")
    cols = _month_columns()
    _band(ws, sp["band"]["A2"], sp["band"]["A3"], stamp, 17)
    sheet = LabelSheet(ws, cols, sp["widths"])
    sheet.header(6, "Portfolio")
    rows = sp["rows"]
    # portfolio block = rows before the first section; off-taker template = first section .. next section (without trailing blanks)
    first_section = next(i for i, x in enumerate(rows) if x["kind"] == "section")
    portfolio_rows = rows[:first_section]
    nxt = next((i for i, x in enumerate(rows) if x["kind"] == "section" and i > first_section), len(rows))
    block = rows[first_section:nxt]
    while block and block[-1]["kind"] == "blank":
        block.pop()
    pf = run.pnl.portfolio

    def month_values(T, key):
        if key not in T:
            return [None] * 13
        v = T[key]
        return [v[12], *v[:12]]

    r = _render_rows(sheet, portfolio_rows, 7, lambda k, row: month_values(pf, k))
    for i, (code, T) in enumerate(run.pnl.sections.items()):
        o = P.offtaker(code)
        title = f"{code} · {o.label}" if o.label != code else code  # display name in the section header (D-C; names never in the repo)
        rows_i = block if i == 0 else [{"kind": "blank"}, *block]  # the portfolio block already ends with its separator
        r = _render_rows(sheet, rows_i, r, lambda k, row, T=T: month_values(T, k), section_text=lambda row, t=title: t)
    _finish(ws, sp["widths"], "C7")
    ws.flush(wb)


def _cf_month(wb: Workbook, run: RunResult, stamp: str) -> None:
    sp = spec()["sheets"]["CF_Mth"]
    ws = Sheet("CF_Mth")
    cols = _month_columns(beyond=True)
    _band(ws, sp["band"]["A2"], sp["band"]["A3"], stamp, 18)
    sheet = LabelSheet(ws, cols, sp["widths"])
    sheet.header(6, "Line")
    rows_ = run.cashflow.rows

    def values_of(key, row):
        if key not in rows_:
            return [None] * 14
        v = rows_[key]
        return [v[13], *v[:12], v[12]]

    _render_rows(sheet, sp["rows"], 7, values_of)
    _finish(ws, sp["widths"], "C7")
    ws.flush(wb)


def _guarantees(wb: Workbook, run: RunResult, stamp: str) -> None:
    sp = spec()["sheets"]["Guarantees"]
    P = run.params
    ws = Sheet("Guarantees")
    cols = _month_columns()
    _band(ws, sp["band"]["A2"], sp["band"]["A3"], stamp, 17)
    sheet = LabelSheet(ws, cols, sp["widths"])
    sheet.header(6, "Line")
    pf = run.pnl.portfolio
    rows = sp["rows"]
    own_start = next((i for i, x in enumerate(rows) if x["kind"] == "line" and x.get("role") == "group" and x.get("key") is None), len(rows))
    head = rows[:own_start]

    def month_values(T, key):
        if key not in T:
            return [None] * 13
        v = T[key]
        return [v[12], *v[:12]]

    r = _render_rows(sheet, head, 7, lambda k, row: month_values(pf, k))
    own = {c: T for c, T in run.pnl.sections.items() if "own_guarantee" in T}
    if own:
        sheet.group(r, "Own guarantees (issued to off-takers)")
        r += 1
        sheet.header(r, "Off-taker")
        r += 1
        for code, T in own.items():
            sheet.line(r, P.offtaker(code).label, "EUR", "data", month_values(T, "own_guarantee"))
            r += 1
    _finish(ws, sp["widths"], "C7")
    ws.flush(wb)


def _pricing(wb: Workbook, run: RunResult, stamp: str) -> None:
    sp = spec()["sheets"]["Pricing_<code>"]
    P = run.params
    for code, pr in run.pricing.items():
        o = P.offtaker(code)
        ws = Sheet(f"Pricing_{code}"[:31])
        cols = _month_columns()
        subtitle = "Year column and months (EUR/MWh unless stated)" if not pr.as_cached else "Workbook-cached logic"
        _band(ws, f"Price build-up - {o.label}", subtitle, stamp, 17)
        sheet = LabelSheet(ws, cols, sp["widths"])
        sheet.header(6, "Line")

        def values_of(key, row, pr=pr):
            if key not in pr.year:
                return [None] * 13
            months = pr.months.get(key)
            return [pr.year[key], *(list(months) if months is not None else [None] * 12)]

        r = _render_rows(sheet, sp["rows"], 7, values_of, unit_ctx="pricing")
        if pr.manual:
            sheet.blank(r)
            r += 1
            sheet.section(r, "Manual case")
            r += 1
            for k, v in pr.manual.items():
                sheet.line(r, label(k, leg=RETAIL), unit_of(k, "pricing"), "data", [v])
                r += 1
        _finish(ws, sp["widths"], "C7")
        ws.flush(wb)


def _ledger(wb: Workbook, run: RunResult, stamp: str) -> None:
    sp = spec()["sheets"]["CF_Daily_Ledger"]
    ws = Sheet("CF_Daily_Ledger")
    d = run.cashflow.daily
    cols = list(d.columns) if d is not None else []
    ncols = max(len(cols), 1)
    _band(ws, sp["band"]["A2"], sp["band"]["A3"], stamp, ncols)
    bg, f = ROLE_STYLE["header"]
    for j, key in enumerate(cols, start=1):
        _put(ws, 6, j, label(key, fallback=TOTAL), f=f, bg=bg, al=CENTER_WRAP)
    ws.height(6, ROW_H_WRAP)
    r = 7
    if d is not None:
        units = [unit_of(k) for k in cols]
        nfs = ["date" if k == "date" else ("month" if k == "month" else number_format(u, "daily")) for k, u in zip(cols, units, strict=True)]
        nfs = [NUMBER_FORMATS["date"]["*"] if x == "date" else (NUMBER_FORMATS["month"]["*"] if x == "month" else x) for x in nfs]
        arr = d.to_numpy(dtype=object)
        dates = pd.to_datetime(d["date"]).dt.date.to_numpy() if "date" in d else None
        for i in range(len(d)):
            for j, key in enumerate(cols, start=1):
                v = dates[i] if key == "date" and dates is not None else arr[i, j - 1]
                if key == "month":
                    v = date(run.params.spine_year, int(v), 1)
                if isinstance(v, pd.Timestamp):
                    v = v.date()
                _put(ws, r, j, v, f=font(), bg=DATA_FILL, nf=nfs[j - 1])
            ws.height(r)
            r += 1
        # 366-day frame: the surplus row sits blank (formatted) in non-leap years
        while r < 7 + 366:
            for j in range(1, len(cols) + 1):
                _put(ws, r, j, None, f=font(), bg=DATA_FILL, nf=nfs[j - 1])
            r += 1
    r += 1
    bg, f = ROLE_STYLE["section"]
    _put(ws, r, 1, "Summary", f=f, bg=bg, al=LEFT)
    for j in range(2, 4):
        _put(ws, r, j, None, f=f, bg=bg)
    r += 1
    bg, f = ROLE_STYLE["header"]
    for j, t in enumerate(("Item", "Unit", "Value"), start=1):
        _put(ws, r, j, t, f=f, bg=bg, al=LEFT if j == 1 else CENTER)
    r += 1
    for k, v in run.cashflow.daily_summary.items():
        u = unit_of(k)
        role = "check" if u == "check" else "data"
        rbg, rf = ROLE_STYLE[role]
        _put(ws, r, 1, label(k, fallback=TOTAL), f=font(), al=LEFT)
        _put(ws, r, 2, unit_text(u), f=font())
        _put(ws, r, 3, v, f=rf, bg=rbg, nf=number_format(u, "daily"))
        r += 1
    widths = dict(sp["widths"])
    for j in range(1, len(cols) + 1):
        widths.setdefault(get_column_letter(j), 25.71)
    _finish(ws, widths, "D7")
    ws.flush(wb)


# ---- wide grids ---------------------------------------------------------------------------------------
KPI_ROWS = ("Total", "Average ( < > 0)", "Maximum", "Minimum")


def _grid_columns(frame: pd.DataFrame) -> list[dict]:
    """Column plan of the wide grids: the calendar block, the sequence, then the engine blocks of the spec with their
    spacers and titles. Engine keys the run does not carry are skipped."""
    sp = spec()["sheets"]["QH_full"]
    have = set(frame.columns)
    plan: list[dict] = [{"kind": "calendar", "key": name, "label": name, "unit": hint, "title": "CALENDAR" if i == 0 else None}
                        for i, (name, hint) in enumerate(CALENDAR_COLUMNS)]
    plan.append({"kind": "info", "key": "seq", "label": "INFO (not consumed) - QH_Sequence_Index", "unit": "#", "title": None})
    titles = {b["col"]: b["title"] for b in sp["blocks"]}
    prev_engine = False
    for c in sp["columns"]:
        if c["kind"] != "engine" or c["key"] == "seq":
            continue
        title = titles.get(c["col"])
        if title or not prev_engine:
            plan.append({"kind": "spacer"})
        if c["key"] in have:
            plan.append({"kind": "engine", "key": c["key"], "label": c["label"], "unit": unit_of(c["key"]), "title": title})
        prev_engine = True
    # drop a spacer that directly follows another spacer
    out = []
    for p in plan:
        if p["kind"] == "spacer" and out and out[-1]["kind"] == "spacer":
            continue
        out.append(p)
    return out


def _grid_header(ws_rows: list[list], plan: list[dict], frame: pd.DataFrame, granularity: str, band_rows: list[list]) -> None:
    """Rows 5-14 of a wide grid as WriteOnlyCell rows (block titles, KPI band, share, key, unit, parameter name)."""
    n = len(plan)
    ws_rows.extend(band_rows)
    # row 5: block titles (Center Across Selection over the block)
    row5 = []
    for p in plan:
        if p["kind"] == "spacer":
            row5.append(("", None, font(18, True, color=WHITE), CAS, None, False))
        else:
            row5.append((p.get("title") or "", INK, font(18, True, color=WHITE), CAS, None, True))
    ws_rows.append(row5)
    ws_rows.append([("", None, font(), CENTER, number_format(p.get("unit"), granularity) if p["kind"] == "engine" else None, False) for p in plan])  # row 6 KPI (empty)
    numeric = {p["key"]: pd.to_numeric(frame[p["key"]], errors="coerce") for p in plan if p["kind"] == "engine"}

    def kpi(name: str, p: dict):
        if p["kind"] != "engine":
            return None
        s = numeric[p["key"]]
        if name == "Total":
            return float(s.sum()) if p["unit"] not in ("EUR/MWh", "%", "ratio", "check", "sign", "#") else None
        if name.startswith("Average"):
            nz = s[s != 0]
            return float(nz.mean()) if len(nz) else None
        return float(s.max()) if name == "Maximum" else float(s.min())

    for name in KPI_ROWS:
        row = []
        for i, p in enumerate(plan):
            if p["kind"] == "spacer":
                row.append(("", None, font(bold=True, color=WHITE), CENTER, None, False))
            else:
                v = name if i == 0 else kpi(name, p)
                row.append((v, INK, font(bold=True, color=WHITE), CENTER, number_format(p.get("unit"), "kpi") if p["kind"] == "engine" else None, True))
        ws_rows.append(row)
    # row 11: share of the block total (engine EUR / MWh columns only) - left empty, formatted
    ws_rows.append([("", None, font(), CENTER, NUMBER_FORMATS["%"]["*"] if p["kind"] == "engine" else None, False) for p in plan])
    ws_rows.append([(p.get("key") if p["kind"] in ("engine", "info") else "", INDEX_FILL if p["kind"] != "spacer" else None, font(10), CENTER, "@", p["kind"] != "spacer") for p in plan])
    ws_rows.append([("Unit" if i == 0 else (unit_text(p.get("unit")) if p["kind"] == "engine" else ""), DATA_FILL if p["kind"] != "spacer" else None, font(bold=True), CENTER, "@", p["kind"] != "spacer") for i, p in enumerate(plan)])
    ws_rows.append([(p.get("label", ""), NAVY if p["kind"] != "spacer" else None, font(bold=True, color=WHITE), CENTER_WRAP, "@", p["kind"] != "spacer") for p in plan])
    assert n == len(row5)


def _grid_sheet(wb: Workbook, run: RunResult, stamp: str, name: str, frame: pd.DataFrame, granularity: str, title: str, subtitle: str) -> None:
    """QH_full (35.040 rows) or QH_daily (366 rows): header block through the buffer, data rows streamed directly."""
    plan = _grid_columns(frame)
    ws = Sheet(name)
    band = []
    for t, f in ((CONFIDENTIAL, font(11, True, color=NAVY)), (title, font(16, True, color=NAVY)), (subtitle, font(10, color=MUTED)), (stamp, font(10, color=MUTED))):
        band.append([(t, None, f, LEFT, None, False)] + [("", None, font(10, color=MUTED), CENTER, None, False)] * (len(plan) - 1))
    rows: list[list] = []
    _grid_header(rows, plan, frame, granularity, band)
    for r_i, row in enumerate(rows, start=1):
        for c_i, (value, bg, f, al, nf, bordered) in enumerate(row, start=1):
            ws.put(r_i, c_i, value if value != "" else None, f=f, bg=bg, nf=nf, al=al, border=bordered)
        ws.height(r_i, ROW_H_WRAP if r_i == 14 else (ROW_H_KEY if r_i == 12 else ROW_H))
    for i, p in enumerate(plan, start=1):
        ws.widths[get_column_letter(i)] = 5.71 if p["kind"] == "spacer" else (25.71 if p["kind"] in ("calendar", "info") else 36.57)
    ws.freeze = "A15"
    ws.flush(wb)
    wso = wb[name]
    # one probe row (row 15) carries every column's data style through openpyxl; the data rows themselves are
    # streamed as XML by _fast_rows after the save (11 million styled cells are far too slow cell by cell)
    nfs = []
    for p in plan:
        if p["kind"] == "spacer":
            nfs.append(None)
        else:
            u = p.get("unit")
            nfs.append(number_format(u, granularity) if p["kind"] == "engine" else NUMBER_FORMATS.get(u, NUMBER_FORMATS["int"])["*"])
    probe = []
    for p, nf in zip(plan, nfs, strict=True):
        if p["kind"] == "spacer":
            probe.append(None)
            continue
        c = WriteOnlyCell(wso, value=None)
        c.font, c.fill, c.alignment, c.number_format, c.border = font(), fill(DATA_FILL), CENTER, nf, BORDER
        probe.append(c)
    wso.append(probe)
    _PENDING_GRIDS.append((name, plan, frame))


_PENDING_GRIDS: list[tuple[str, list[dict], pd.DataFrame]] = []
_EPOCH = date(1899, 12, 30)


def _xml_escape(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _fast_rows(xlsx: bytes, grids: list[tuple[str, list[dict], pd.DataFrame]]) -> bytes:
    """Replace the probe row of each wide grid by the frame's rows, written straight as sheet XML with the probe's
    style ids (same styles, same workbook - only the cell-by-cell object churn is skipped). The rows are streamed
    into the zip entry so the 400 MB of sheet XML of QH_full never sit in memory."""
    import re
    import zipfile
    from xml.etree import ElementTree as ET

    NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    zin = zipfile.ZipFile(io.BytesIO(xlsx))
    wbx = ET.fromstring(zin.read("xl/workbook.xml"))
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))}
    parts = {s.get("name"): "xl/" + rels[s.get(RID)].replace("/xl/", "") for s in wbx.find(NS + "sheets")}
    by_part = {parts[name]: (plan, frame) for name, plan, frame in grids}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename not in by_part:
                zout.writestr(item, zin.read(item.filename))
                continue
            plan, frame = by_part[item.filename]
            xml = zin.read(item.filename)
            m = re.search(rb'<row r="15"[^>]*>(.*?)</row>', xml, re.S)
            styles = dict(re.findall(rb'<c r="([A-Z]+)15" s="(\d+)"', m.group(1)))
            head, tail = xml[: m.start()], xml[m.end():]
            cols = []
            for i, p in enumerate(plan, start=1):
                if p["kind"] == "spacer":
                    continue
                letter = get_column_letter(i).encode()
                col = frame[p["key"]]
                if p["kind"] == "engine":
                    kind = "num" if pd.api.types.is_numeric_dtype(col) else "text"
                else:
                    hint = p.get("unit")
                    kind = {"int": "num", "date": "date", "month": "date", "time": "time", "#": "num"}.get(hint, "text")
                cols.append((letter, styles.get(letter, b"0"), kind, col.to_numpy()))
            ht = b' ht="%s" customHeight="1"' % str(ROW_H).encode()
            with zout.open(item.filename, "w") as fh:
                fh.write(head)
                chunk: list[bytes] = []
                for i in range(len(frame)):
                    rb = str(15 + i).encode()
                    cells = []
                    for letter, sid, kind, arr in cols:
                        v = arr[i]
                        if v is None or (isinstance(v, float) and v != v):
                            cells.append(b'<c r="%s%s" s="%s"/>' % (letter, rb, sid))
                        elif kind == "num":
                            cells.append(b'<c r="%s%s" s="%s"><v>%s</v></c>' % (letter, rb, sid, (str(int(v)) if isinstance(v, (int, np.integer)) else repr(float(v))).encode()))
                        elif kind == "date":
                            d = v if isinstance(v, date) else pd.Timestamp(v).date()
                            cells.append(b'<c r="%s%s" s="%s"><v>%d</v></c>' % (letter, rb, sid, (d - _EPOCH).days))
                        elif kind == "time":
                            cells.append(b'<c r="%s%s" s="%s"><v>%s</v></c>' % (letter, rb, sid, repr((v.hour * 3600 + v.minute * 60 + v.second) / 86400).encode()))
                        else:
                            cells.append(b'<c r="%s%s" s="%s" t="inlineStr"><is><t>%s</t></is></c>' % (letter, rb, sid, _xml_escape(str(v)).encode("utf-8")))
                    chunk.append(b'<row r="%s"%s>%s</row>' % (rb, ht, b"".join(cells)))
                    if len(chunk) >= 256:
                        fh.write(b"".join(chunk))
                        chunk.clear()
                fh.write(b"".join(chunk))
                fh.write(tail)
    return buf.getvalue()


def _qh_frames(run: RunResult) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(full, daily): calendar block + seq + engine keys at quarter-hour, and the daily aggregate (means for prices)."""
    q = run.qh.qh
    cal = calendar_block(run.grid)
    full = pd.concat([cal, q.drop(columns=[c for c in ("date", "interval", "month", "peak") if c in q.columns]).reset_index(drop=True)], axis=1)
    eng = [c for c in q.columns if q[c].dtype.kind == "f"]
    means = {c for c in eng if unit_of(c) in ("EUR/MWh", "%", "ratio", "sign")}
    daily = q.groupby(run.grid["date"].dt.date.values)[eng].agg({c: ("mean" if c in means else "sum") for c in eng})
    day_cal = cal.groupby(cal["Date_EET"], sort=True).first().loc[daily.index]
    day_cal.index.name = "Date_EET"
    daily = pd.concat([day_cal.reset_index(), daily.reset_index(drop=True), pd.DataFrame({"seq": np.arange(1, len(daily) + 1)})], axis=1)
    return full.copy(), daily.copy()


def qh_csv_frame(run: RunResult, full: bool = True) -> pd.DataFrame:
    """The quarter-hour (or daily) frame in the one column order of the workbook (D-G): the calendar block, the
    sequence, then the engine keys by block - no spacers. This is the CSV download of the Engine page."""
    q_full, q_daily = _qh_frames(run)
    frame = q_full if full else q_daily
    plan = _grid_columns(frame)
    cols = [p["key"] for p in plan if p["kind"] != "spacer"]
    return frame[cols]


QH_CSV_KEYS_NOTE = "calendar block, sequence, then the engine keys in the order of the workbook's QH sheets (D-G)"


# ---- parameters and provenance -----------------------------------------------------------------------
PARAM_FORMATS = {"EUR/MWh": '#,##0.00\\ "€/MWh";\\(#,##0.00\\ "€/MWh"\\);"-"', "EUR": _nf(0, "€"), "RON": _nf(0, "RON"), "%": "0.0%", "% p.a.": '0.00%',
                 "days": "0", "day of month": "0", "months": "0", "MW": '#,##0.0\\ "MW"', "RON/MW": '#,##0\\ "RON/MW"', "RON/MWh": '#,##0.00\\ "RON/MWh"',
                 "RON/EUR": NUMBER_FORMATS["RON/EUR"]["*"], "RON/GC": NUMBER_FORMATS["RON/GC"]["*"], "GC/MWh": NUMBER_FORMATS["GC/MWh"]["*"],
                 "x": "0.00", "date": NUMBER_FORMATS["date"]["*"], "year": "0", "sign": "0", "#": "0"}


def _param_value(v, unit: str):
    """The register value as the cell should carry it: dates as dates, numbers as numbers, lists and text as text."""
    if isinstance(v, str) and unit == "date":
        try:
            return date.fromisoformat(v)
        except ValueError:
            return v
    if isinstance(v, bool) or v is None:
        return "" if v is None else ("yes" if v else "no")
    if isinstance(v, (int, float)):
        return v
    return str(v)


def _parameters(wb: Workbook, run: RunResult, stamp: str, scenario_name: str) -> None:
    from esb.catalogue import catalogue_for

    sp = spec()["sheets"]["Parameters"]
    ws = Sheet("Parameters")
    _band(ws, sp["band"]["A2"], f"Scenario file: {scenario_name or '(unsaved)'}", stamp, 5)
    hf = font(11, True, color=WHITE)
    for j, t in enumerate(("Parameter", "Value", "Unit", "Standard / default", "Source / vintage"), start=1):
        _put(ws, 6, j, t, f=hf, bg=NAVY, al=LEFT if j in (1, 4, 5) else CENTER)
    ws.height(6)
    flat = _flatten(run.params.to_dict())
    catalogue = catalogue_for(flat.keys())
    src, status = holiday_source()
    r = 7
    for k, v in flat.items():
        meta = catalogue[k]
        val = _param_value(v, meta.unit)
        nf = PARAM_FORMATS.get(meta.unit, "General") if not isinstance(val, str) else "@"
        _put(ws, r, 1, k, f=font(10), al=LEFT)
        _put(ws, r, 2, val, f=font(10), bg=DATA_FILL, nf=nf)
        _put(ws, r, 3, meta.unit, f=font(10), bg=DATA_FILL)
        _put(ws, r, 4, meta.standard, f=font(10), bg=DATA_FILL, al=LEFT)
        _put(ws, r, 5, meta.source_text, f=font(10), bg=DATA_FILL, al=LEFT)
        ws.height(r)
        r += 1
    _put(ws, r, 1, "calendar.public_holidays", f=font(10), al=LEFT)
    _put(ws, r, 2, "rules in config/calendar_ro.yaml", f=font(10), bg=DATA_FILL, nf="@")
    _put(ws, r, 3, "", f=font(10), bg=DATA_FILL)
    _put(ws, r, 4, "Codul muncii art. 139 alin. (1)", f=font(10), bg=DATA_FILL, al=LEFT)
    _put(ws, r, 5, f"{src} [{status}]", f=font(10), bg=DATA_FILL, al=LEFT)
    ws.height(r)
    _finish(ws, sp["widths"], "B7")
    ws.flush(wb)


def _provenance(wb: Workbook, run: RunResult, stamp: str, provenance: list[dict] | None, requested_by: str) -> None:
    sp = spec()["sheets"]["Provenance"]
    ws = Sheet("Provenance")
    headers = list(sp["header"].values())
    _band(ws, sp["band"]["A2"], sp["band"]["A3"], stamp, len(headers))
    hf = font(11, True, color=WHITE)
    for j, t in enumerate(headers, start=1):
        _put(ws, 6, j, t, f=hf, bg=NAVY, al=LEFT)
    ws.height(6)
    r = 7
    for rec in provenance or []:
        row = _provenance_row(rec, requested_by)
        for j, v in enumerate(row, start=1):
            _put(ws, r, j, v, f=font(10), al=LEFT_WRAP)
        ws.height(r)
        r += 1
    r += 1
    bg, f = ROLE_STYLE["section"]
    _put(ws, r, 1, "Calculation-order trace (G0-D6)", f=f, bg=bg, al=LEFT)
    for j in range(2, 4):
        _put(ws, r, j, None, f=f, bg=bg)
    r += 1
    for j, t in enumerate(("Stage", "Unit", "Cumulative"), start=1):
        _put(ws, r, j, t, f=hf, bg=NAVY, al=LEFT if j == 1 else CENTER)
    r += 1
    for stage, secs in run.trace:
        _put(ws, r, 1, stage, f=font(10), al=LEFT)
        _put(ws, r, 2, "s", f=font(10))
        _put(ws, r, 3, secs, f=font(10), bg=DATA_FILL, nf="0.000")
        r += 1
    _finish(ws, sp["widths"], "B7")
    ws.flush(wb)


def _provenance_row(rec: dict, requested_by: str) -> list:
    """Map an assembled-series source record (esb.assemble) to the 11-column provenance schema of the template:
    Written to | Source system | Object | Obtained via | Query or request | Parameters | Run at | Requested by |
    Row count | Changes after retrieval | Notes."""
    kind = rec.get("kind", "")
    if kind == "reference":
        return ["series (all classes)", "Reference Case fixture", rec.get("detail", ""), "bundled with the application", "", "", "", requested_by, "", "none", ""]
    params = "; ".join(x for x in (f"scenario {rec['scenario']}" if rec.get("scenario") else "", f"time basis {rec.get('time_basis', '')}",
                                   f"k {rec['k']}" if rec.get("k") else "") if x)
    notes = ("not delivered: " + ", ".join(rec["not_delivered"])) if rec.get("not_delivered") else ""
    return [f"series ({kind})", "upload (ESB-STD-QH template)", rec.get("filename", ""), "Data page upload", f"md5 {rec.get('md5', '')}", params,
            rec.get("imported_at_utc", ""), requested_by, rec.get("rows", ""), "none - accepted bytes are kept byte-identical in the case bundle", notes]


# ---- roles for the application (D-H) ------------------------------------------------------------------
@cache
def roles_for(sheet: str, block: str | None = None) -> dict[str, str]:
    """{engine key: role} of a spec sheet (optionally one block: overview, overview_cf, overview_checks, portfolio,
    section, cashflow, pricing, ledger_summary), so the application shows the same roles as the workbook."""
    sp = spec()["sheets"].get(sheet, {})
    out: dict[str, str] = {}
    for row in sp.get("rows", []) + sp.get("summary", []):
        if row.get("kind") == "line" and row.get("key") and (block is None or row.get("block") == block):
            out.setdefault(row["key"], row.get("role", "data"))
    return out


# ---- entry point ----------------------------------------------------------------------------------------
SHEET_ORDER = ("Portf Overview", "Cons_P&L", "CF_Mth", "CF_Daily_Ledger", "Pricing_<code>", "Guarantees", "QH_daily", "QH_full", "Parameters", "Provenance")


def build_workbook(run: RunResult, include_qh: bool = False, scenario_name: str = "", provenance: list[dict] | None = None,
                   requested_by: str = "") -> bytes:
    """The canonical export: Portf Overview, Cons_P&L, CF_Mth, CF_Daily_Ledger, Pricing_<code>, Guarantees, QH_daily,
    QH_full (on request), Parameters, Provenance - rendered from data/layout/run_export.json in write-only mode."""
    stamp = _stamp(run, scenario_name)
    _PENDING_GRIDS.clear()
    wb = Workbook(write_only=True)
    wb._fonts = IndexedList([font()])  # the workbook default (Normal style) is Montserrat 12 - untouched cells inherit it
    wb._named_styles["Normal"].font = font()
    _overview(wb, run, stamp)
    _cons_pnl(wb, run, stamp)
    _cf_month(wb, run, stamp)
    _ledger(wb, run, stamp)
    _pricing(wb, run, stamp)
    _guarantees(wb, run, stamp)
    full, daily = _qh_frames(run)
    _grid_sheet(wb, run, stamp, "QH_daily", daily, "daily", "QH P&L - daily", "Day sums of the engine keys (means for prices, shares and ratios); calendar block first")
    if include_qh:
        _grid_sheet(wb, run, stamp, "QH_full", full, "qh", "QH P&L", "Every quarter-hour of the spine year: calendar block, then the engine keys by block")
    _parameters(wb, run, stamp, scenario_name)
    _provenance(wb, run, stamp, provenance, requested_by)
    buf = io.BytesIO()
    wb.save(buf)
    grids, _PENDING_GRIDS[:] = list(_PENDING_GRIDS), []
    return _fast_rows(buf.getvalue(), grids)
