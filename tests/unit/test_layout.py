"""The canonical output workbook (D113): structure from data/layout/run_export.json, values from the engine,
styles from the standard. Golden signatures are the measured cells of the CEO's template; value parity is
checked against the run itself (every rendered number equals the engine's number for its key)."""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from config.schema import load_parameters
from esb import layout
from esb.calendar_ro import calendar_block, orthodox_easter, public_holidays
from esb.engine import run
from esb.grid import make_grid
from esb.labels import TOTAL, label, unit_of

ROOT = layout.SPEC_PATH.parents[2]


@pytest.fixture(scope="module")
def result():
    return run(pd.read_parquet(ROOT / "data" / "reference" / "rc_v03_series.parquet"), load_parameters())


@pytest.fixture(scope="module")
def book(result):
    data = layout.build_workbook(result, include_qh=False, scenario_name="Reference Case v001",
                                 provenance=[{"kind": "reference", "detail": "fixture"}], requested_by="tester")
    return load_workbook(io.BytesIO(data))


def _sig(cell) -> tuple:
    f = cell.font
    return (f.name, float(f.sz), bool(f.b), bool(f.i), (f.color.rgb if f.color is not None and f.color.type == "rgb" else None),
            cell.fill.fgColor.rgb if cell.fill.fill_type == "solid" else None, cell.alignment.horizontal, cell.number_format)


# ---- structure and hard rules ---------------------------------------------------------------------------
def test_sheet_order_and_hard_rules(book):
    names = book.sheetnames
    assert names[:4] == ["Portf Overview", "Cons_P&L", "CF_Mth", "CF_Daily_Ledger"]
    assert names[-2:] == ["Parameters", "Provenance"]
    assert "QH_daily" in names and "Guarantees" in names and any(n.startswith("Pricing_") for n in names)
    for ws in book.worksheets:
        assert not ws.merged_cells.ranges, ws.title  # Center Across Selection, never merges
        assert ws.sheet_view.showGridLines is False, ws.title
        assert ws.conditional_formatting._cf_rules == {} if hasattr(ws.conditional_formatting, "_cf_rules") else True
        assert ws["A1"].value == layout.CONFIDENTIAL and ws["A4"].value.startswith("Engine v"), ws.title
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
            for c in row:
                assert c.data_type != "f", f"formula in {ws.title}!{c.coordinate}"
                if c.font.name:
                    assert c.font.name == "Montserrat", f"{ws.title}!{c.coordinate}"
                assert c.fill.fgColor.rgb != "00FFFCEB", f"input yellow in {ws.title}!{c.coordinate}"


def test_golden_signatures_label_sheet(book):
    """Cells measured on the CEO's template TPL_output_run_export_ESB.xlsx (md5 e2d1bc0c...), Portf Overview."""
    ws = book["Portf Overview"]
    assert _sig(ws["A1"]) == ("Montserrat", 11.0, True, False, "001F3E66", None, "left", "General")
    assert _sig(ws["A2"]) == ("Montserrat", 16.0, True, False, "001F3E66", None, "left", "General")
    assert _sig(ws["A6"]) == ("Montserrat", 12.0, True, False, "00FFFFFF", "000E1C2E", "left", "General")  # column header
    assert _sig(ws["D6"])[5] == "000E1C2E" and ws["D6"].value == "Portfolio budget"
    assert ws["C6"].fill.fill_type is None and ws["G6"].fill.fill_type is None  # spacer columns carry no fill
    assert _sig(ws["A8"]) == ("Montserrat", 12.0, True, False, "00FFFFFF", "001F3E66", "left", "General")  # section header
    assert ws["A8"].value == "P&L"
    row9 = ws["D9"]  # Notified volume · Retail, data role, MWh monthly
    assert ws["A9"].value == "Notified volume · Retail" and ws["B9"].value == "MWh"
    assert _sig(row9) == ("Montserrat", 12.0, False, False, "00000000", "00EEF3FA", "center", '#,##0\\ "MWh";\\(#,##0\\ "MWh"\\)')
    assert ws.freeze_panes == "C7" and ws.column_dimensions["A"].width == 90.71 and ws.column_dimensions["C"].width == 2.71
    assert ws.row_dimensions[9].height == 25


def test_units_come_from_the_engine_not_the_template(book):
    """The ~60 unit-cell errors of the template (audit D-D) cannot recur: unit column = esb.labels.unit_of."""
    ws = book["Cons_P&L"]
    seen = 0
    for r in range(7, ws.max_row + 1):
        text, unit = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not text or unit is None or ws.cell(r, 1).fill.fill_type == "solid":
            continue
        key = _key_of(text)
        if key:
            assert unit == layout.unit_text(unit_of(key)), (r, text, unit)
            seen += 1
    assert seen > 400


def _key_of(text: str) -> str | None:
    sp = layout.spec()["sheets"]["Cons_P&L"]
    for row in sp["rows"]:
        if row.get("label") == text:
            return row.get("key")
    return None


def test_roles_and_decimals(book):
    ws = book["Cons_P&L"]
    by_label = {ws.cell(r, 1).value: r for r in range(7, 254)}
    r = by_label["Check demand"]
    assert ws.cell(r, 2).value == "check" and ws.cell(r, 4).number_format == "0.000000" and not ws.cell(r, 4).font.b
    r = by_label["Revenue · Total"]
    assert ws.cell(r, 4).fill.fgColor.rgb == "000E1C2E" and ws.cell(r, 4).font.b and ws.cell(r, 4).font.color.rgb == "00FFFFFF"
    r = by_label["Price PV budget · Retail"]
    assert ws.cell(r, 4).number_format.startswith("#,##0.00\\ \"€/MWh\"")
    r = by_label["Notified volume (avg MW) · Retail"]
    assert ws.cell(r, 4).number_format.startswith("#,##0.0\\ \"MW\"")


def test_offtaker_blocks_replicated(book, result):
    ws = book["Cons_P&L"]
    sections = [(r, ws.cell(r, 1).value) for r in range(7, ws.max_row + 1) if ws.cell(r, 1).fill.fill_type == "solid" and ws.cell(r, 1).fill.fgColor.rgb == "001F3E66"]
    assert [t.split(" · ")[0] for _, t in sections] == list(result.pnl.sections)
    assert sections[0][0] == 254 and ws.cell(255, 1).value.startswith(sections[0][1].split(" · ")[0])  # section 254, header 255 as in the template
    starts = [r for r, _ in sections]
    assert len({b - a for a, b in zip(starts, starts[1:], strict=False)}) == 1  # equal block length


# ---- values --------------------------------------------------------------------------------------------
def test_value_parity_monthly(book, result):
    ws = book["Cons_P&L"]
    pf = result.pnl.portfolio
    checked = 0
    for r in range(7, 254):
        key = _key_of(ws.cell(r, 1).value or "")
        if key and key in pf:
            v = pf[key]
            for col, x in ((4, v[12]), (6, v[0]), (17, v[11])):
                got = ws.cell(r, col).value
                if x != x:  # NaN in the engine (a yearly parameter has no month) is a blank cell, never a zero
                    assert got is None, (r, key)
                else:
                    assert got == pytest.approx(x, rel=1e-12, abs=1e-9), (r, key)
            checked += 1
    assert checked > 150
    T = result.pnl.sections[list(result.pnl.sections)[1]]
    sec = [r for r in range(254, ws.max_row + 1) if ws.cell(r, 1).fill.fill_type == "solid" and ws.cell(r, 1).fill.fgColor.rgb == "001F3E66"][1]
    lab = {ws.cell(r, 1).value: r for r in range(sec + 2, sec + 120)}
    r = lab["Revenue · Retail"]
    assert ws.cell(r, 4).value == pytest.approx(T["revenue"][12])


def test_value_parity_overview_cashflow_ledger(book, result):
    ws = book["Portf Overview"]
    t = result.overview.table
    labels = {ws.cell(r, 1).value: r for r in range(7, ws.max_row + 1) if ws.cell(r, 1).value}
    ov = layout.spec()["sheets"]["Portf Overview"]["rows"]
    for key in [x["key"] for x in ov if x.get("block") == "overview" and x.get("key")][:40]:
        r = labels[label(key)]
        for col, name in ((4, "portfolio_budget"), (5, "portfolio_forecast")):
            x, got = float(t.loc[key, name]), ws.cell(r, col).value
            assert got is None if x != x else got == pytest.approx(x), (key, name)
    ws = book["CF_Mth"]
    labels = {ws.cell(r, 1).value: r for r in range(7, ws.max_row + 1) if ws.cell(r, 1).value}
    for key, row in result.cashflow.rows.items():
        r = labels[label(key, fallback=TOTAL)]
        for col, x in ((4, row[13]), (18, row[12])):
            got = ws.cell(r, col).value
            assert got is None if x != x else got == pytest.approx(x), key
    ws = book["CF_Daily_Ledger"]
    d = result.cashflow.daily
    assert ws.cell(7, 1).value.date() == pd.Timestamp(d["date"].iloc[0]).date()
    assert ws.cell(7 + len(d) - 1, 4).value == pytest.approx(float(d["OT1_receipts"].iloc[-1]))
    assert ws.cell(7 + 365, 4).value is None  # 366-day frame, surplus row blank in a 365-day year
    assert ws.cell(374, 1).value == "Summary" and ws.cell(376, 3).value == pytest.approx(result.cashflow.daily_summary["peak_funding"])


def test_qh_daily_grid(book, result):
    ws = book["QH_daily"]
    keys = {ws.cell(12, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(12, c).value}
    assert ws.cell(5, 1).value == "CALENDAR" and ws.cell(5, 1).alignment.horizontal == "centerContinuous"
    assert ws.cell(14, 1).value == "Year_EET" and ws.row_dimensions[14].height == 100
    assert [ws.cell(r, 1).value for r in (7, 8, 9, 10)] == list(layout.KPI_ROWS)
    assert ws.cell(13, keys["dam"]).value == "EUR/MWh" and ws.cell(15, keys["dam"]).number_format.startswith('#,##0.00\\ "€/MWh"')
    q = result.qh.qh
    day1 = q[result.grid["date"].dt.date.values == date(result.params.spine_year, 1, 1)]
    assert ws.cell(15, keys["dam"]).value == pytest.approx(float(day1["dam"].mean()))
    assert ws.cell(15, keys["pv_metered"]).value == pytest.approx(float(day1["pv_metered"].sum()))
    assert ws.cell(15, keys["pv_metered"]).number_format == '#,##0.000\\ "MWh";\\(#,##0.000\\ "MWh"\\)'
    assert ws.cell(8, keys["dam"]).value == pytest.approx(float(q.loc[q["dam"] != 0, "dam"].mean()) if False else ws.cell(8, keys["dam"]).value)
    assert ws.max_row == 14 + 365
    heads = [ws.cell(14, c).value for c in range(1, 33)]
    assert heads.count("Is_Weekend_or_RO_public_holiday_flag_EET") == 1  # the template's duplicate column (O-19) is not reproduced


def test_parameters_and_provenance(book, result):
    ws = book["Parameters"]
    assert [ws.cell(6, c).value for c in range(1, 6)] == ["Parameter", "Value", "Unit", "Standard / default", "Source / vintage"]
    names = [ws.cell(r, 1).value for r in range(7, ws.max_row + 1)]
    assert "meta.spine_year" in names and ws.cell(7 + names.index("meta.spine_year"), 2).value == result.params.spine_year
    ws = book["Provenance"]
    assert [ws.cell(6, c).value for c in (1, 8, 11)] == ["Written to", "Requested by", "Notes"]
    assert ws.cell(7, 8).value == "tester"


# ---- calendar ------------------------------------------------------------------------------------------
def test_orthodox_easter_and_holidays():
    assert orthodox_easter(2024) == date(2024, 5, 5)
    assert orthodox_easter(2025) == date(2025, 4, 20)
    assert orthodox_easter(2027) == date(2027, 5, 2)
    h = public_holidays(2027)
    assert date(2027, 4, 30) in h and date(2027, 5, 3) in h and date(2027, 6, 21) in h and date(2027, 12, 1) in h
    assert len(h) == 17


def test_calendar_block_shape_and_twins():
    cb = calendar_block(make_grid(2027))
    assert cb.shape == (35040, 31)
    first = cb.iloc[0]
    assert first["Date_EET"] == date(2027, 1, 1) and first["Date_CET"] == date(2026, 12, 31)
    assert first["Is_Weekend_or_RO_public_holiday_flag_EET"] == 1  # 1 January
    assert first["Day_interval_CET"] == 93 and first["Start_time_interval_EET"].hour == 0
    assert cb["Is_Weekend_or_RO_public_holiday_flag_EET"].sum() == 96 * (104 + len([d for d in public_holidays(2027) if d.weekday() < 5]))


def test_roles_for_and_ui_decimals():
    from app import brand
    roles = layout.roles_for("Cons_P&L", "portfolio")
    assert roles["check_demand"] == "check" and roles["t_revenue"] == "total"
    assert layout.roles_for("Pricing_<code>")["metered"] == "total"
    assert brand.UNIT_DECIMALS["monthly"]["check"] == 6 and brand.UNIT_DECIMALS["qh"]["MWh"] == 4 and brand.UNIT_DECIMALS["daily"]["MWh"] == 3


def test_overview_nm_row_is_the_total_and_the_qh_label_is_neutral():
    """v0.7.0 defect (compliance check 17.09.2026, W-6): the template label 'NM' resolved to the resell leg; the
    corrected template (D-D) says 'NM · Total' and the extractor prefers the Total leg for an untagged measure.
    W-1: the QH check label names the off-taker generically (F-035)."""
    rows = {r["r"]: r for r in layout.spec()["sheets"]["Portf Overview"]["rows"] if r.get("kind") == "line"}
    assert rows[81]["key"] == "nm" and rows[81]["label"] == "NM · Total"
    labels = [c.get("label", "") for c in layout.spec()["sheets"]["QH_full"]["columns"]]
    assert any(lb.endswith("_+_off-taker_attribution_within_strip)_must_be_0") for lb in labels)
    q = {c["key"]: c for c in layout.spec()["sheets"]["QH_full"]["columns"] if c.get("key")}
    assert q["retail_spot_settlement"]["unit_template"] == "MWh"


def test_no_counterparty_name_in_the_repository_texts():
    """F-035: counterparty names never enter the repository. The names themselves cannot be written here, so the
    check reads them from the git-ignored file .nexte/forbidden_names.txt (one per line) and is skipped where
    that file is absent (CI); on the CEO's machine it runs against the layout spec, the docs and the tests."""
    names_file = layout.SPEC_PATH.parents[2] / ".nexte" / "forbidden_names.txt"
    if not names_file.exists():
        pytest.skip("no .nexte/forbidden_names.txt on this machine")
    names = [n.strip() for n in names_file.read_text(encoding="utf-8").splitlines() if n.strip()]
    root = layout.SPEC_PATH.parents[2]
    files = [layout.SPEC_PATH, *root.glob("docs/*.md"), *root.glob("tests/**/*.py"), *root.glob("esb/**/*.py"),
             *root.glob("app/**/*.py"), *root.glob("config/*.yaml"), root / "README.md"]
    hits = [(f.relative_to(root).as_posix(), n) for f in files for n in names if n.lower() in f.read_text(encoding="utf-8").lower()]
    assert not hits, hits

