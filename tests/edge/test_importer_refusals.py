"""End-to-end: template -> filled workbook -> importer, including every refusal path."""

from datetime import date

import numpy as np
import pytest
from openpyxl import load_workbook

from esb import grid
from esb.importer import build_template, import_workbook, write_delivery
from esb.importer.contract import Control, Registry, SlotSpec
from tests.conftest import YEAR, local_clock_frame, synthetic_registry


@pytest.fixture(scope="module")
def good_workbook(tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / "good.xlsx"
    control = Control(input_class="wholesale_prices", spine_year=YEAR, time_basis="local_clock", provider="test")
    write_delivery(path, control, synthetic_registry(), local_clock_frame())
    return path


def _check(result, code):
    return next(c for c in result.checks if c.code == code)


def test_template_builds_for_every_class(tmp_path):
    for cls in ("offtaker_load", "pv_generation", "baseload_nomination", "wholesale_prices"):
        p = build_template(cls, YEAR, tmp_path / f"{cls}.xlsx")
        wb = load_workbook(p, read_only=True)
        assert set(["Instructions", "Std_Control", "Series_Registry", "RAW_EET_QH", "Recon_Check"]) <= set(wb.sheetnames)
        wb.close()
    # an empty template is refused (no data), not crashed
    r = import_workbook(tmp_path / "offtaker_load.xlsx")
    assert not r.ok and not _check(r, "2").passed


def test_good_local_clock_workbook_is_accepted(good_workbook):
    r = import_workbook(good_workbook)
    assert r.ok, r.summary()
    assert all(c.passed for c in r.checks) and len(r.checks) == 6
    f = r.frame
    assert len(f) == 365 * 96
    raw = local_clock_frame()
    assert abs(f["vol_MWh"].sum() - raw["vol_MWh"].sum()) < 1e-6
    assert r.provenance.k == {"vol_MWh": -1}
    assert r.provenance.md5 and r.provenance.input_class == "wholesale_prices"
    spring, _ = grid.dst_dates(YEAR)
    inj = f[(f["date"].dt.date == spring) & f["interval"].between(13, 16)]
    assert (inj["vol_MWh"] == 0).all() and np.allclose(inj["price_EUR_MWh"], [113, 114, 115, 116])


def test_formula_in_raw_is_refused(good_workbook, tmp_path):
    p = tmp_path / "formula.xlsx"
    wb = load_workbook(good_workbook)
    wb["RAW_EET_QH"]["D100"] = "=D99+1"
    wb.save(p)
    r = import_workbook(p)
    assert not r.ok and not _check(r, "3").passed
    assert "formula" in _check(r, "3").detail


def test_undeclared_column_with_data_is_refused(good_workbook, tmp_path):
    p = tmp_path / "undeclared.xlsx"
    wb = load_workbook(good_workbook)
    ws = wb["RAW_EET_QH"]
    ws["H1"] = "E07"
    ws["H2"] = 1.0
    wb.save(p)
    r = import_workbook(p)
    assert not r.ok and "undeclared" in _check(r, "3").detail


def test_blank_in_required_series_is_refused_and_never_filled(good_workbook, tmp_path):
    p = tmp_path / "blank.xlsx"
    wb = load_workbook(good_workbook)
    wb["RAW_EET_QH"]["D5000"] = None
    wb.save(p)
    r = import_workbook(p)
    assert not r.ok and not _check(r, "6").passed and "vol_MWh" in _check(r, "6").detail


def test_blank_in_optional_series_is_reported_not_refused(good_workbook, tmp_path):
    p = tmp_path / "blank_opt.xlsx"
    wb = load_workbook(good_workbook)
    wb["RAW_EET_QH"]["G5000"] = None  # C01 direction, optional
    wb.save(p)
    r = import_workbook(p)
    assert r.ok and "optional series with blanks: direction (1)" in _check(r, "6").detail


def test_missing_rows_are_refused(good_workbook, tmp_path):
    p = tmp_path / "missing.xlsx"
    wb = load_workbook(good_workbook)
    wb["RAW_EET_QH"].delete_rows(300, 3)
    wb.save(p)
    r = import_workbook(p)
    assert not r.ok and not _check(r, "2").passed and "3 missing" in _check(r, "2").detail


def test_fixed_96_file_declared_as_local_clock_is_refused(tmp_path):
    g = grid.make_grid(YEAR)
    import pandas as pd

    raw = pd.DataFrame(
        {
            "Date_EET": g["date"].dt.date.values,
            "Start_EET": [t.time() for t in g["start_eet"]],
            "End_EET": [t.time() for t in g["end_eet"]],
            "vol_MWh": 1.0,
            "price_EUR_MWh": 50.0,
            "ratio": 0.0,
            "direction": "Balanced",
        }
    )
    p = tmp_path / "wrong_basis.xlsx"
    write_delivery(p, Control(input_class="wholesale_prices", spine_year=YEAR, time_basis="local_clock"), synthetic_registry(), raw)
    r = import_workbook(p)
    assert not r.ok and "4 missing, 4 duplicated" in _check(r, "2").detail
    # the same rows declared fixed_96 are accepted
    p2 = tmp_path / "right_basis.xlsx"
    write_delivery(p2, Control(input_class="wholesale_prices", spine_year=YEAR, time_basis="fixed_96"), synthetic_registry(), raw)
    assert import_workbook(p2).ok


def test_volume_series_without_k_is_refused(tmp_path):
    reg = Registry([SlotSpec("E01", "vol_MWh", "MWh", "Extensive", "inject_zero", "sum", basis="metered")])
    with pytest.raises(ValueError, match="k = -1"):
        write_delivery(tmp_path / "x.xlsx", Control("offtaker_load", YEAR), reg, None)


def test_rows_outside_the_spine_year_are_refused(good_workbook, tmp_path):
    p = tmp_path / "year.xlsx"
    wb = load_workbook(good_workbook)
    wb["RAW_EET_QH"]["A2"] = date(YEAR - 1, 12, 31)
    wb.save(p)
    r = import_workbook(p)
    assert not r.ok and "outside the spine year" in "; ".join(_check(r, "2").items)


def test_wrong_control_block_is_refused(good_workbook, tmp_path):
    p = tmp_path / "control.xlsx"
    wb = load_workbook(good_workbook)
    ws = wb["Std_Control"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "intervals_per_day":
            row[1].value = 24
    wb.save(p)
    r = import_workbook(p)
    assert not r.ok and not _check(r, "1").passed
